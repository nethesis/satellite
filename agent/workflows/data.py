"""Bounded table ingestion and deterministic identity/payment lookup."""

import base64
import csv
import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation
from urllib.parse import quote, urlsplit, urljoin, parse_qs

from agent.application.contracts import ApplicationError, fields, identifier, integer, string

MAX_UPLOAD = 10 * 1024**2
MAX_ROWS = 50000
MAX_CELLS = 1000000
CANONICAL_FIELDS = {"resident_id", "name", "phone", "building", "unit", "period", "amount", "currency", "verification_code", "customer_id"}
REQUIRED = {"resident_id", "name", "period", "amount", "currency"}


def normalized_name(value):
    return " ".join(re.sub(r"[^\w\s]", " ", "".join(c for c in unicodedata.normalize("NFKD", str(value)).casefold() if not unicodedata.combining(c))).split())


def matching_name(actual, expected, mode='exact'):
    if actual == expected:
        return True
    if mode != 'similar':
        return False
    actual = actual.replace(' ', ''); expected = expected.replace(' ', '')
    if actual == expected:
        return True
    if min(len(actual), len(expected)) < 4 or abs(len(actual)-len(expected)) > 1:
        return False
    if len(actual) > len(expected):
        actual, expected = expected, actual
    left = right = edits = 0
    while left < len(actual) and right < len(expected):
        if actual[left] == expected[right]:
            left += 1; right += 1
        else:
            edits += 1
            if edits > 1:
                return False
            if len(actual) == len(expected):
                left += 1
            right += 1
    return edits + len(actual)-left + len(expected)-right <= 1


def normalized_phone(value, country_code="39"):
    value = str(value or "").strip()
    if not value or re.search(r"[A-Za-z]", value):
        return ""
    digits = re.sub(r"\D", "", value)
    if value.startswith("00"):
        digits = digits[2:]
    elif not value.startswith("+") and not (digits.startswith(country_code) and len(digits) > 10):
        digits = country_code + digits
    return "+" + digits if 6 <= len(digits) <= 15 else ""


def settings(value):
    fields(value, ("name", "format", "mapping", "country_code", "decimal_separator", "header_row"),
           ("delimiter", "sheet", "text_pattern", "secret_ref", "spreadsheet_id", "range", "refresh_seconds", "published_url"))
    string(value["name"], 128)
    if value["format"] not in ("csv", "xlsx", "text", "google_sheets", "google_csv"):
        raise ApplicationError("unsupported_data_format")
    if not isinstance(value["mapping"], dict) or not REQUIRED <= value["mapping"].keys() or not value["mapping"].keys() <= CANONICAL_FIELDS:
        raise ApplicationError("invalid_column_mapping")
    for column in value["mapping"].values():
        string(column, 128)
    if not re.fullmatch(r"[0-9]{1,4}", value["country_code"]) or value["decimal_separator"] not in (".", ","):
        raise ApplicationError("invalid_locale")
    integer(value["header_row"], 1, 100)
    if "delimiter" in value and (not isinstance(value["delimiter"], str) or len(value["delimiter"]) != 1 or value["delimiter"] in "\r\n\0"):
        raise ApplicationError("invalid_delimiter")
    if value["format"] == "google_sheets":
        identifier(value.get("secret_ref"))
        if not re.fullmatch(r"[A-Za-z0-9_-]{10,128}", value.get("spreadsheet_id", "")):
            raise ApplicationError("invalid_spreadsheet_id")
        string(value.get("range"), 128)
    if value['format'] == 'google_csv':
        published_csv_url(value.get('published_url'), initial=True)
    if value['format'] in ('google_sheets', 'google_csv'):
        integer(value.get("refresh_seconds", 0), 0, 86400)
        if 0 < value.get("refresh_seconds", 0) < 300:
            raise ApplicationError("refresh_too_frequent")
    if "text_pattern" in value:
        string(value["text_pattern"], 2048)
    return value


def published_csv_url(value, *, initial=False):
    string(value, 2048)
    url = urlsplit(value)
    try:
        allowed = url.scheme == 'https' and url.port in (None, 443) and not url.username and not url.password and not url.fragment
    except ValueError:
        allowed = False
    host = url.hostname or ''
    if not allowed or not (host == 'docs.google.com' or not initial and re.fullmatch(r'doc-[a-z0-9-]+-sheets\.googleusercontent\.com', host)):
        raise ApplicationError('invalid_published_sheet_url')
    if initial:
        query = parse_qs(url.query, keep_blank_values=True)
        if not re.fullmatch(r'/spreadsheets/d/e/2PACX-[A-Za-z0-9_-]{40,256}/pub', url.path) or set(query) - {'gid', 'single', 'output'} or query.get('output') != ['csv'] or query.get('single') != ['true'] or len(query.get('gid', [])) != 1 or not re.fullmatch(r'[0-9]{1,10}', query['gid'][0]):
            raise ApplicationError('invalid_published_sheet_url')
    return host


async def published_csv(cfg):
    """Fetch only a Google-published CSV; redirects retain public-address/TLS checks."""
    import asyncio
    import aiohttp
    from agent.application.http import PolicyResolver
    url = cfg['published_url']
    published_csv_url(url, initial=True)
    try:
        async with asyncio.timeout(15):
            for attempt in range(4):
                host = published_csv_url(url)
                connector = aiohttp.TCPConnector(resolver=PolicyResolver(host, []), use_dns_cache=False, ssl=True)
                async with aiohttp.ClientSession(connector=connector, trust_env=False) as session:
                    async with session.get(url, allow_redirects=False, headers={'Accept':'text/csv','Accept-Encoding':'identity'}) as response:
                        if response.status in (301, 302, 303, 307, 308):
                            url = urljoin(url, response.headers.get('Location', ''))
                            published_csv_url(url)
                            continue
                        if response.status != 200 or response.content_type not in ('text/csv', 'application/octet-stream', 'text/plain'):
                            raise ApplicationError('published_sheet_unavailable', 503)
                        raw = bytearray()
                        async for chunk in response.content.iter_chunked(65536):
                            raw.extend(chunk)
                            if len(raw) > MAX_UPLOAD:
                                raise ApplicationError('upload_too_large', 413)
                        return bytes(raw)
        raise ApplicationError('published_sheet_redirect_limit', 503)
    except ApplicationError:
        raise
    except (aiohttp.ClientError, OSError, TimeoutError):
        raise ApplicationError('published_sheet_unavailable', 503) from None


def table_values(raw, cfg):
    if len(raw) > MAX_UPLOAD:
        raise ApplicationError("upload_too_large", 413)
    if cfg["format"] == "xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 64 * 1024**2:
                    raise ApplicationError("workbook_too_large")
                if any("vbaProject" in name for name in archive.namelist()):
                    raise ApplicationError("macros_not_supported")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook[cfg.get("sheet") or workbook.sheetnames[0]]
                if sheet.max_row > MAX_ROWS + 100 or sheet.max_column > 128:
                    raise ApplicationError("workbook_too_large")
                result = []
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ApplicationError("formula_values_require_export")
                    result.append([cell.value for cell in cells])
                return result
            finally:
                workbook.close()
        except ApplicationError:
            raise
        except Exception:
            raise ApplicationError("invalid_workbook") from None
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ApplicationError("invalid_text_encoding") from None
    if cfg["format"] == "text" and cfg.get("text_pattern"):
        import regex
        try:
            pattern = regex.compile(cfg["text_pattern"])
            columns = list(cfg["mapping"].values())
            rows = [columns]
            for line in text.splitlines():
                if not line.strip():
                    continue
                match = pattern.fullmatch(line, timeout=0.02)
                if not match:
                    raise ApplicationError("unparsed_text_line")
                rows.append([match.group(column) for column in columns])
                if len(rows) > MAX_ROWS + 1:
                    raise ApplicationError("too_many_rows")
            return rows
        except ApplicationError:
            raise
        except Exception:
            raise ApplicationError("invalid_or_slow_text_pattern") from None
    reader = csv.reader(io.StringIO(text), delimiter=cfg.get("delimiter", ","))
    rows = []
    for row in reader:
        if len(row) > 128 or len(rows) >= MAX_ROWS + 100:
            raise ApplicationError("table_too_large")
        rows.append(row)
    return rows


def normalize_table(values, cfg):
    header_index = 0 if cfg["format"] == "text" and cfg.get("text_pattern") else cfg["header_row"] - 1
    if len(values) <= header_index:
        raise ApplicationError("missing_header")
    header = [str(x or "").strip() for x in values[header_index]]
    if len(set(header)) != len(header) or not set(cfg["mapping"].values()) <= set(header):
        raise ApplicationError("missing_or_duplicate_columns")
    positions = {field: header.index(column) for field, column in cfg["mapping"].items()}
    result = []; keys = set(); identities = {}; errors = []; cell_count = 0
    for index, cells in enumerate(values[header_index + 1:], header_index + 2):
        cell_count += len(cells)
        if len(result) >= MAX_ROWS or cell_count > MAX_CELLS:
            raise ApplicationError("table_too_large")
        if not any(str(x or "").strip() for x in cells):
            continue
        try:
            row = {field: str(cells[position] if position < len(cells) and cells[position] is not None else "").strip() for field, position in positions.items()}
            if any(len(value) > 512 for value in row.values()) or any(not row[field] for field in REQUIRED):
                raise ValueError("missing_value")
            if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", row["period"]) or not re.fullmatch(r"[A-Z]{3}", row["currency"]):
                raise ValueError("invalid_period_or_currency")
            amount_text = row["amount"]
            # An explicit locale avoids interpreting thousands/decimal separators heuristically.
            if cfg["decimal_separator"] == ",":
                amount_text = amount_text.replace(".", "").replace(",", ".")
            elif "," in amount_text:
                raise ValueError("invalid_decimal_separator")
            amount = Decimal(amount_text)
            if not amount.is_finite() or amount < 0 or amount > Decimal("10000000") or amount != amount.quantize(Decimal("0.01")):
                raise ValueError("invalid_amount")
            row["amount"] = format(amount.quantize(Decimal("0.01")), "f")
            row["phone"] = normalized_phone(row.get("phone"), cfg["country_code"])
            row["name_key"] = normalized_name(row["name"])
            row["source_row"] = index
            key = (row["resident_id"], row["period"])
            if key in keys:
                raise ValueError("duplicate_resident_period")
            evidence = (row['name_key'], row['phone'], row.get('verification_code', ''))
            if row['resident_id'] in identities and identities[row['resident_id']] != evidence:
                raise ValueError('inconsistent_resident_identity')
            identities[row['resident_id']] = evidence
            keys.add(key); result.append(row)
        except (ValueError, InvalidOperation) as exc:
            errors.append({"row": index, "error": str(exc) if isinstance(exc, ValueError) else "invalid_amount"})
            if len(errors) >= 100:
                break
    if not result and not errors:
        raise ApplicationError("empty_table")
    return result, errors


def ingest(raw, cfg):
    return normalize_table(table_values(raw, cfg), cfg)


def _parse_worker(connection, raw, cfg):
    try:
        connection.send((True, ingest(raw, cfg)))
    except ApplicationError as exc:
        connection.send((False, exc.code))
    except Exception:
        connection.send((False, "invalid_data_file"))
    finally:
        connection.close()


async def bounded_ingest(raw, cfg):
    """A timed-out parser is terminated, rather than left running in a thread."""
    import asyncio
    import multiprocessing
    mp = multiprocessing.get_context("spawn")
    receiver, sender = mp.Pipe(duplex=False)
    worker = mp.Process(target=_parse_worker, args=(sender, raw, cfg), daemon=True)
    worker.start(); sender.close()
    try:
        success, value = await asyncio.wait_for(asyncio.to_thread(receiver.recv), 15)
        if not success:
            raise ApplicationError(value)
        return value
    finally:
        if worker.is_alive():
            worker.terminate()
        await asyncio.to_thread(worker.join, 2)
        receiver.close()


def safe_row(row):
    return {key: value for key, value in row.items() if key not in ("verification_code", "name_key", "phone")}


async def google_values(cfg, credential, transport):
    import json
    from functools import partial
    import asyncio
    from google.oauth2.service_account import Credentials
    from google.auth.transport.requests import Request
    import requests
    info = json.loads(credential)
    if info.get("type") != "service_account" or info.get("token_uri") != "https://oauth2.googleapis.com/token":
        raise ApplicationError("invalid_service_account")
    credentials = Credentials.from_service_account_info(info, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    session = requests.Session(); session.trust_env = False
    try:
        await asyncio.to_thread(credentials.refresh, partial(Request(session=session), timeout=15))
    finally:
        session.close()
    path = "/v4/spreadsheets/" + quote(cfg["spreadsheet_id"], safe="") + "/values/" + quote(cfg["range"], safe="")
    result = await transport.json_request("https://sheets.googleapis.com", path, "GET",
        {"Authorization": "Bearer " + credentials.token}, {"valueRenderOption": "UNFORMATTED_VALUE"}, None, [], 15, max_bytes=MAX_UPLOAD)
    values = result.get("values")
    if not isinstance(values, list) or len(values) > MAX_ROWS + 100 or any(not isinstance(row, list) or len(row) > 128 for row in values):
        raise ApplicationError("invalid_sheet_data")
    return values
